"""把 collect.ps1 采集到的 CSV 做成可翻阅的 Excel 工作簿。"""
import csv
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "app-capabilities.csv")
OUT = os.path.join(HERE, "安卓应用能力清单.xlsx")

CAPABILITIES = [
    "接收分享·文字", "接收分享·图片", "接收分享·视频", "接收分享·任意文件", "接收分享·多个文件",
    "处理选中文字", "编辑图片", "打开网页链接", "打开PDF", "打开视频文件", "打开音频文件",
    "当文件选择器", "拨号", "发短信", "发邮件", "地图导航", "建日程", "设闹钟", "设计时器",
    "建联系人", "拍照", "录音", "语音识别", "建备忘", "网页搜索", "扫码", "设壁纸", "文字朗读(TTS)",
]

HEAD_FILL = PatternFill("solid", fgColor="2E6BE6")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(size=10)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
OK_FONT = Font(size=10, color="1B7F3B", bold=True)
OK_FILL = PatternFill("solid", fgColor="E8F6EC")
MONO = Font(size=9, name="Consolas")


def style_header(ws, row=1, ncols=None):
    ncols = ncols or ws.max_column
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 30


def read_rows():
    with open(SRC, encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def sheet_matrix(wb, rows):
    ws = wb.active
    ws.title = "能力矩阵"
    headers = ["名称", "包名", "类型", "系统角色", "能力数", "公开URL协议"] + CAPABILITIES
    ws.append(headers)

    rows = sorted(rows, key=lambda r: (-int(r["能力数"] or 0), r["名称"] or r["包名"]))
    for r in rows:
        caps = set((r["能力"] or "").split())
        ws.append(
            [
                r["名称"] or "",
                r["包名"],
                r["类型"],
                r["系统角色"] or "",
                int(r["能力数"] or 0),
                r["公开协议"] or "",
            ]
            + ["✅" if c in caps else "" for c in CAPABILITIES]
        )

    style_header(ws)
    widths = {"A": 20, "B": 42, "C": 8, "D": 16, "E": 7, "F": 40}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for i in range(len(CAPABILITIES)):
        ws.column_dimensions[get_column_letter(7 + i)].width = 11

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.border = BORDER
            cell.font = BODY_FONT
            if cell.column == 2:
                cell.font = MONO
            if cell.column >= 7:
                cell.alignment = Alignment(horizontal="center")
                if cell.value == "✅":
                    cell.font = OK_FONT
                    cell.fill = OK_FILL
            elif cell.column in (3, 5):
                cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    return ws


def sheet_index(wb, rows):
    ws = wb.create_sheet("能力反向索引")
    ws.append(["能力", "能做到的 App 数", "都有谁"])
    for cap in CAPABILITIES:
        who = [
            (r["名称"] or r["包名"])
            for r in rows
            if cap in set((r["能力"] or "").split())
        ]
        ws.append([cap, len(who), "、".join(who) if who else "（无）"])
    style_header(ws)
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 110
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=3):
        for cell in row:
            cell.border = BORDER
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=(cell.column == 3))
        row[1].alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    return ws


SCHEME_TESTS = [
    ("weixin://", "微信", "实测：只有微信认领"),
    ("taobao://", "淘宝", "实测"),
    ("alipays://", "支付宝", "实测"),
    ("zhihu://", "知乎", "实测"),
    ("github://", "GitHub 官方 App", "实测；它还认领 https://github.com 与 https://*.ghe.com"),
    ("qq://", "小爱同学", "实测：QQ 没认领，被小爱同学认领了"),
    ("bilibili://", "（无人响应）", "实测：B站装了但没注册这个协议"),
    ("xhsdiscover://", "（无人响应）", "实测：小红书用的是别的协议"),
    ("https://github.com/...", "小米浏览器、GitHub App、UC", "实测：同一个链接有 3 条路，Agent 可以选"),
    ("https://（通用）", "小米浏览器、UC", "实测：装了 3 个浏览器，夸克没认领通用 https"),
    ("smsto:10086", "短信", "发短信的唯一入口"),
    ("mailto:a@b.com", "邮件、Gmail", "发邮件的两个入口"),
    ("geo:0,0?q=...", "高德地图、百度地图", "导航的两个入口"),
    ("tel:", "联系人", "拨号盘"),
    ("content://com.android.contacts/contacts", "联系人", "建联系人"),
    ("vnd.android.cursor.item/event", "日历", "建日程"),
]


def sheet_schemes(wb):
    ws = wb.create_sheet("常见URL协议实测")
    ws.append(["接口 / 协议", "这台手机上谁能响应", "备注"])
    for row in SCHEME_TESTS:
        ws.append(list(row))
    style_header(ws)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 60
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=3):
        for cell in row:
            cell.border = BORDER
            cell.font = MONO if cell.column == 1 else BODY_FONT
            if cell.column == 3:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    return ws


NOTES = [
    ("这份表是什么", ""),
    ("", "手机（小米 2509FPN0BC，Android 16 / SDK 36）上每个 App 对外暴露的能力清单。"),
    ("", "一行一个 App，一列一个它能对外提供的能力。只列出至少能干一件事的 App。"),
    ("", ""),
    ("怎么采集的", ""),
    ("", "1) cmd package query-activities / query-services —— 问系统「谁能处理这个请求」。"),
    ("", "   只返回导出的、真正可被调用的组件，所以这是「能不能用」的准确来源。"),
    ("", "2) dumpsys package 的 * Resolver Table -> Schemes —— 每个 App 注册了哪些 URL 协议。"),
    ("", "3) pm list packages -s / -3 —— 区分系统应用与第三方应用。"),
    ("", "重跑：pwsh -NoProfile -ExecutionPolicy Bypass -File collect.ps1"),
    ("", ""),
    ("必须知道的边界", ""),
    ("", "· 枚举只能发现 App「已经开好的门」。App 没声明的接口，永远枚举不出来。"),
    ("", "  例：微信只声明了接收分享（SEND），没有任何「发送消息」接口 —— 实测 query SENDTO 给微信返回空。"),
    ("", "· 绝大多数 URL 协议是推送/支付/统计 SDK 的内部管道，不是给人用的入口。"),
    ("", "  「公开URL协议」列已滤掉一批明显的垃圾，原始全量在 app-capabilities.csv 的 URL协议全量 列。"),
    ("", "· 这份数据是 adb（系统身份）看到的，比普通 App 能看到的更多。"),
    ("", "  Android 11+ 起，App 必须逐个声明 <queries> 才能看见对应的 App，"),
    ("", "  或者申请敏感的 QUERY_ALL_PACKAGES（Play 需申报用途）。"),
    ("", "· 「建备忘」= 0：安卓没有系统级备忘录接口。手机里装了小米笔记，但它不认标准动作。"),
]


def sheet_notes(wb):
    ws = wb.create_sheet("说明与方法")
    ws.append(["项目", "内容"])
    for a, b in NOTES:
        ws.append([a, b])
    style_header(ws)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 105
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=2):
        for cell in row:
            cell.border = BORDER
            cell.font = BODY_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if row[0].value:
            row[0].font = Font(size=10, bold=True)
    return ws


def main():
    rows = read_rows()
    wb = Workbook()
    sheet_matrix(wb, rows)
    sheet_index(wb, rows)
    sheet_schemes(wb)
    sheet_notes(wb)
    wb.save(OUT)
    print(f"rows={len(rows)} -> {OUT}")


if __name__ == "__main__":
    main()
